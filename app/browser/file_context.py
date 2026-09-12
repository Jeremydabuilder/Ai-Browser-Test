"""Local file parsing for context the user explicitly attaches.

Every caller of ``read_local_file`` already has a path the user themselves
picked - a native file dialog, not a path the model chose. That is the
whole safety story here: this module has no discovery, no directory
listing, and nothing that takes a path from model output. It only ever
turns bytes the user pointed at into text.

Supported: .txt, .md, .json, .csv, .docx, .pdf (via app.browser.pdf_context).
Anything else raises FileParsingError rather than guessing at a format.
"""

from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from pathlib import Path

from app.browser.pdf_context import PdfExtractionError, extract_pdf

#: A file larger than this is refused outright, before any parsing is
#: attempted - the point is to fail fast and say why, not to let a huge
#: upload silently blow past MAX_TEXT_CHARS after doing all the work anyway.
MAX_FILE_BYTES = 20 * 1024 * 1024
#: Extracted text is capped the same way page text and PDF text are - the
#: model's own context budget would truncate it anyway.
MAX_TEXT_CHARS = 40000

_PLAIN_TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
SUPPORTED_SUFFIXES = _PLAIN_TEXT_SUFFIXES | {".json", ".csv", ".docx", ".pdf"}


class FileParsingError(RuntimeError):
    """Always carries a message safe to show the person who chose the file."""


@dataclass(frozen=True)
class FileDocument:
    source: str      # the path the user chose
    filename: str
    kind: str         # "text", "json", "csv", "docx", "pdf"
    text: str
    truncated: bool


def _read_bytes(path: Path) -> bytes:
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise FileParsingError(f"Could not read that file: {exc}") from exc
    if size > MAX_FILE_BYTES:
        raise FileParsingError(
            f"That file is too large ({size // (1024 * 1024)} MB). "
            f"The limit is {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    try:
        return path.read_bytes()
    except OSError as exc:
        raise FileParsingError(f"Could not read that file: {exc}") from exc


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FileParsingError("That file does not look like readable text.")


def _parse_csv(data: bytes) -> str:
    text = _decode_text(data)
    try:
        rows = list(csv.reader(io.StringIO(text)))
    except csv.Error as exc:
        raise FileParsingError(f"That doesn't look like a valid CSV file: {exc}") from exc
    return "\n".join(", ".join(cell for cell in row) for row in rows)


def _parse_json(data: bytes) -> str:
    text = _decode_text(data)
    try:
        parsed = json.loads(text)
    except ValueError as exc:
        raise FileParsingError(f"That doesn't look like valid JSON: {exc}") from exc
    return json.dumps(parsed, indent=2, ensure_ascii=False)


def _parse_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise FileParsingError(
            "Reading .docx files needs the 'python-docx' package, which is "
            "not installed.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx's own errors vary by failure mode
        raise FileParsingError(f"That doesn't look like a valid .docx file: {exc}") from exc
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs)


def read_local_file(path: str) -> FileDocument:
    """Parse a file the user explicitly picked. Raises FileParsingError for
    an unsupported extension, an unreadable path, or a file too large or
    too malformed to make sense of."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise FileParsingError(
            f"'{suffix or file_path.name}' is not a supported file type. "
            f"Supported: {', '.join(sorted(SUPPORTED_SUFFIXES))}.")

    if suffix == ".pdf":
        data = _read_bytes(file_path)
        try:
            document = extract_pdf(f"file://{file_path}", data=data)
        except PdfExtractionError as exc:
            raise FileParsingError(str(exc)) from exc
        text = document.full_text
        kind = "pdf"
    else:
        data = _read_bytes(file_path)
        if suffix == ".csv":
            text, kind = _parse_csv(data), "csv"
        elif suffix == ".json":
            text, kind = _parse_json(data), "json"
        elif suffix == ".docx":
            text, kind = _parse_docx(data), "docx"
        else:
            text, kind = _decode_text(data), "text"

    truncated = len(text) > MAX_TEXT_CHARS
    return FileDocument(source=str(file_path), filename=file_path.name, kind=kind,
                        text=text[:MAX_TEXT_CHARS], truncated=truncated)


__all__ = ["FileDocument", "FileParsingError", "SUPPORTED_SUFFIXES", "read_local_file"]
